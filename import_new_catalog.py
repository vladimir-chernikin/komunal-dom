#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
ИМПОРТ НОВОГО КАТАЛОГА УСЛУГ ИЗ EXCEL

Excel файл: /obmen/cat.xlsx
Лист: "Каталог услуг"
Структура: 44 услуги

Использование:
    python import_new_catalog.py
"""

import os
import sys
import django
from pathlib import Path

# Setup Django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'komunal_dom.settings')
os.environ.setdefault('DJANGO_ALLOW_ASYNC_UNSAFE', 'true')
sys.path.insert(0, '/var/www/komunal-dom_ru')
django.setup()

import openpyxl
from django.db import connection


def import_catalog_from_excel():
    """Импортирует 44 услуги из Excel в services_catalog_v2"""

    excel_path = Path('/obmen/cat.xlsx')

    if not excel_path.exists():
        print(f"❌ Файл не найден: {excel_path}")
        return False

    print(f"✅ Файл найден: {excel_path}")

    # Загружаем Excel
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb['Каталог услуг']

    print(f"✅ Лист: {ws.title}")
    print(f"✅ Размер: {ws.max_row - 1} услуг (без заголовка)")

    # Создаем таблицу services_catalog_v2
    with connection.cursor() as cursor:
        # Проверяем, существует ли таблица
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'services_catalog_v2'
            );
        """)
        exists = cursor.fetchone()[0]

        if exists:
            print("⚠️  Таблица services_catalog_v2 уже существует - удаляем")
            cursor.execute("DROP TABLE services_catalog_v2;")

        # Создаем таблицу services_catalog_v2
        print("📝 Создаем таблицу services_catalog_v2...")
        cursor.execute("""
            CREATE TABLE services_catalog_v2 (
                service_id SERIAL PRIMARY KEY,
                scenario_name VARCHAR(255) NOT NULL,
                type_name VARCHAR(100) NOT NULL,
                localization_name VARCHAR(100) NOT NULL,
                category_name VARCHAR(255) NOT NULL,
                description TEXT,
                route_name VARCHAR(150),
                is_internal BOOLEAN DEFAULT FALSE,
                is_active BOOLEAN DEFAULT TRUE,
                created_at TIMESTAMP WITH TIME ZONE DEFAULT NOW(),
                updated_at TIMESTAMP WITH TIME ZONE DEFAULT NOW()
            );
        """)

        # Создаем индексы
        cursor.execute("CREATE INDEX idx_services_catalog_v2_category ON services_catalog_v2(category_name);")
        cursor.execute("CREATE INDEX idx_services_catalog_v2_type ON services_catalog_v2(type_name);")
        cursor.execute("CREATE INDEX idx_services_catalog_v2_localization ON services_catalog_v2(localization_name);")
        cursor.execute("CREATE INDEX idx_services_catalog_v2_active ON services_catalog_v2(is_active);")

        print("✅ Таблица services_catalog_v2 создана")

    # Импортируем данные
    print(f"\n📥 Импортируем услуги...")

    imported = 0
    skipped = 0

    with connection.cursor() as cursor:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:  # Пропускаем пустые строки
                skipped += 1
                continue

            (
                service_id_excel,
                scenario_name,
                type_name,
                localization_name,
                category_name,
                description,
                route_name,
                is_internal
            ) = row

            # Преобразуем is_internal
            is_internal_bool = (is_internal == 'Да')

            # Вставляем запись
            cursor.execute("""
                INSERT INTO services_catalog_v2 (
                    scenario_name,
                    type_name,
                    localization_name,
                    category_name,
                    description,
                    route_name,
                    is_internal,
                    is_active
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            """, (
                scenario_name,
                type_name,
                localization_name,
                category_name,
                description,
                route_name,
                is_internal_bool,
                True
            ))

            imported += 1

            if imported <= 5:
                print(f"  [{imported}] {scenario_name[:50]}...")

    print(f"\n✅ Импортировано: {imported} услуг")
    print(f"⏭️  Пропущено: {skipped} пустых строк")

    # Проверяем результат
    with connection.cursor() as cursor:
        cursor.execute("SELECT COUNT(*) FROM services_catalog_v2;")
        count = cursor.fetchone()[0]
        print(f"📊 Всего в services_catalog_v2: {count} записей")

        # Показываем примеры
        cursor.execute("""
            SELECT service_id, scenario_name, type_name, localization_name, category_name, is_internal
            FROM services_catalog_v2
            ORDER BY service_id
            LIMIT 10;
        """)
        print(f"\n📋 Первые 10 услуг:")
        for row in cursor.fetchall():
            (sid, name, typ, loc, cat, internal) = row
            internal_mark = "🔧" if internal else "✨"
            print(f"  [{sid}] {name[:40]}... | {typ} | {loc} | {cat} {internal_mark}")

    return True


if __name__ == '__main__':
    print("="*80)
    print("ИМПОРТ НОВОГО КАТАЛОГА УСЛУГ")
    print("="*80)
    print()

    try:
        success = import_catalog_from_excel()
        if success:
            print("\n" + "="*80)
            print("✅ ИМПОРТ ЗАВЕРШЕН УСПЕШНО")
            print("="*80)
        else:
            print("\n" + "="*80)
            print("❌ ИМПОРТ ЗАВЕРШЕН С ОШИБКАМИ")
            print("="*80)
            sys.exit(1)
    except Exception as e:
        print(f"\n❌ ОШИБКА: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
